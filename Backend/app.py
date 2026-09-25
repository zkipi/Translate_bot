from flask import Flask, request, jsonify, send_file
from flask_cors import CORS

from translator import translate_text

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

import pandas as pd

from io import BytesIO

import threading
import uuid


# ==========================================
# SETUP
# ==========================================

app = Flask(__name__)

CORS(app)


# ==========================================
# LANGUAGE COLUMN SUFFIXES
# ==========================================

LANGUAGE_SUFFIXES = {
    "English": "en",
    "Swedish": "sv",
    "Danish": "da",
    "Norwegian": "no",
    "Finnish": "fi"
}


# ==========================================
# EXCEL JOB STORAGE
# ==========================================

excel_jobs = {}


# ==========================================
# EXCEL FORMATTING
# ==========================================

def format_worksheet(worksheet):

    # ------------------------------------------
    # COLORS
    # ------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    alternate_fill = PatternFill(
        fill_type="solid",
        fgColor="F2F6FA"
    )


    # ------------------------------------------
    # HEADER
    # ------------------------------------------

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )


    for cell in worksheet[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # ------------------------------------------
    # BODY
    # ------------------------------------------

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        # Alternate row background

        if row_number % 2 == 0:

            for cell in worksheet[row_number]:

                cell.fill = alternate_fill


        # Cell formatting

        for cell in worksheet[row_number]:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


    # ------------------------------------------
    # COLUMN WIDTHS
    # ------------------------------------------

    for column in worksheet.columns:

        column_letter = (
            column[0].column_letter
        )

        header = worksheet[
            f"{column_letter}1"
        ].value


        # Description columns

        if (
            header == "description"
            or (
                isinstance(header, str)
                and header.startswith(
                    "description_"
                )
            )
        ):

            worksheet.column_dimensions[
                column_letter
            ].width = 60


        # Product name columns

        elif header in [
            "name",
            "title",
            "product_name"
        ]:

            worksheet.column_dimensions[
                column_letter
            ].width = 35


        # ID columns

        elif header in [
            "id",
            "ID"
        ]:

            worksheet.column_dimensions[
                column_letter
            ].width = 12


        # Other columns

        else:

            worksheet.column_dimensions[
                column_letter
            ].width = 20


    # ------------------------------------------
    # ROW HEIGHT
    # ------------------------------------------

    worksheet.row_dimensions[1].height = 25


    for row_number in range(
        2,
        worksheet.max_row + 1
    ):

        worksheet.row_dimensions[
            row_number
        ].height = 80


    # ------------------------------------------
    # BORDERS
    # ------------------------------------------

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D9E1F2"
        )
    )


    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.border = thin_border


    # ------------------------------------------
    # FREEZE HEADER
    # ------------------------------------------

    worksheet.freeze_panes = "A2"


    # ------------------------------------------
    # FILTER
    # ------------------------------------------

    if worksheet.max_row >= 2:

        worksheet.auto_filter.ref = (
            worksheet.dimensions
        )


# ==========================================
# HOME
# ==========================================

@app.route("/")
def home():

    return "Translate Bot is running!"


# ==========================================
# TRANSLATE TEXT
# ==========================================

@app.route(
    "/translate",
    methods=["POST"]
)
def translate():

    data = request.get_json()

    text = data.get("text")

    target_language = data.get(
        "target_language"
    )


    # ------------------------------------------
    # CHECK INPUT
    # ------------------------------------------

    if not text:

        return jsonify({
            "error": "No text provided"
        }), 400


    if not target_language:

        return jsonify({
            "error":
                "No target language provided"
        }), 400


    # ------------------------------------------
    # TRANSLATE
    # ------------------------------------------

    try:

        result = translate_text(
            text,
            target_language
        )


        return jsonify({

            "translation":
                result["translation"],

            "source_language":
                result["source_language"],

            "target_language":
                target_language

        })


    except Exception as error:

        print(
            "Translation error:",
            error
        )


        return jsonify({
            "error":
                "Translation failed"
        }), 500


# ==========================================
# EXPORT SINGLE TRANSLATION
# ==========================================

@app.route(
    "/export-excel",
    methods=["POST"]
)
def export_excel():

    data = request.get_json()


    original_text = data.get(
        "original_text"
    )

    translation = data.get(
        "translation"
    )

    source_language = data.get(
        "source_language"
    )

    target_language = data.get(
        "target_language"
    )


    # ------------------------------------------
    # CHECK INPUT
    # ------------------------------------------

    if not original_text:

        return jsonify({
            "error":
                "No original text provided"
        }), 400


    if not translation:

        return jsonify({
            "error":
                "No translation provided"
        }), 400


    # ------------------------------------------
    # CREATE WORKBOOK
    # ------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Translation"


    # ------------------------------------------
    # HEADERS
    # ------------------------------------------

    worksheet.append([
        "Original description",
        "Translation",
        "Source language",
        "Target language"
    ])


    # ------------------------------------------
    # DATA
    # ------------------------------------------

    worksheet.append([
        original_text,
        translation,
        source_language,
        target_language
    ])


    # ------------------------------------------
    # FORMAT
    # ------------------------------------------

    format_worksheet(
        worksheet
    )


    # ------------------------------------------
    # SAVE
    # ------------------------------------------

    excel_file = BytesIO()

    workbook.save(
        excel_file
    )

    excel_file.seek(0)


    # ------------------------------------------
    # SEND FILE
    # ------------------------------------------

    return send_file(
        excel_file,

        as_attachment=True,

        download_name="translated_product.xlsx",

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# ==========================================
# START EXCEL TRANSLATION JOB
# ==========================================

@app.route(
    "/translate-excel",
    methods=["POST"]
)
def translate_excel():

    # ------------------------------------------
    # CHECK FILE
    # ------------------------------------------

    if "file" not in request.files:

        return jsonify({
            "error":
                "No Excel file provided"
        }), 400


    file = request.files["file"]


    target_language = request.form.get(
        "target_language"
    )


    # ------------------------------------------
    # CHECK LANGUAGE
    # ------------------------------------------

    if not target_language:

        return jsonify({
            "error":
                "No target language provided"
        }), 400


    # ------------------------------------------
    # CHECK FILE
    # ------------------------------------------

    if not file.filename:

        return jsonify({
            "error":
                "No file selected"
        }), 400


    # ------------------------------------------
    # READ EXCEL
    # ------------------------------------------

    try:

        dataframe = pd.read_excel(
            file
        )


    except Exception as error:

        print(
            "Excel read error:",
            error
        )


        return jsonify({
            "error":
                "Could not read the Excel file."
        }), 400


    # ------------------------------------------
    # CHECK DESCRIPTION COLUMN
    # ------------------------------------------

    if "description" not in dataframe.columns:

        return jsonify({
            "error":
                'The Excel file must contain a '
                'column named "description".'
        }), 400


    # ------------------------------------------
    # CREATE JOB
    # ------------------------------------------

    job_id = str(
        uuid.uuid4()
    )


    total_rows = len(
        dataframe
    )


    excel_jobs[job_id] = {

        "status": "processing",

        "current": 0,

        "total": total_rows,

        "completed": 0,

        "failed": 0,

        "download_ready": False,

        "file": None,

        "error": None

    }


    # ------------------------------------------
    # START BACKGROUND THREAD
    # ------------------------------------------

    thread = threading.Thread(
        target=process_excel_translation,

        args=(
            job_id,
            dataframe,
            target_language
        )
    )


    thread.start()


    # ------------------------------------------
    # RETURN JOB ID
    # ------------------------------------------

    return jsonify({

        "job_id": job_id,

        "total": total_rows

    })


# ==========================================
# PROCESS EXCEL TRANSLATION
# ==========================================

def process_excel_translation(
    job_id,
    dataframe,
    target_language
):

    try:

        # --------------------------------------
        # TARGET COLUMN
        # --------------------------------------

        suffix = LANGUAGE_SUFFIXES.get(
            target_language,
            target_language.lower()
        )


        target_column = (
            f"description_{suffix}"
        )


        translations = []


        total_rows = len(
            dataframe
        )


        print(
            f"Starting Excel translation "
            f"job {job_id}"
        )


        # --------------------------------------
        # TRANSLATE ROWS
        # --------------------------------------

        for index, value in enumerate(
            dataframe["description"]
        ):

            # ----------------------------------
            # EMPTY CELL
            # ----------------------------------

            if pd.isna(value):

                translations.append("")

                excel_jobs[job_id][
                    "current"
                ] = index + 1

                excel_jobs[job_id][
                    "completed"
                ] += 1

                print(
                    f"Skipped empty row "
                    f"{index + 1}/"
                    f"{total_rows}"
                )

                continue


            # ----------------------------------
            # CONVERT TO TEXT
            # ----------------------------------

            text = str(value).strip()


            # ----------------------------------
            # EMPTY STRING
            # ----------------------------------

            if not text:

                translations.append("")

                excel_jobs[job_id][
                    "current"
                ] = index + 1

                excel_jobs[job_id][
                    "completed"
                ] += 1

                continue


            # ----------------------------------
            # PUNCTUATION ONLY
            # ----------------------------------

            if not any(
                character.isalnum()
                for character in text
            ):

                translations.append(
                    text
                )

                excel_jobs[job_id][
                    "current"
                ] = index + 1

                excel_jobs[job_id][
                    "completed"
                ] += 1

                print(
                    f"Skipped row "
                    f"{index + 1}: "
                    f"no translatable text"
                )

                continue


            # ----------------------------------
            # TRANSLATE
            # ----------------------------------

            try:

                result = translate_text(
                    text,
                    target_language
                )


                translation = (
                    result["translation"]
                )


                translations.append(
                    translation
                )


                excel_jobs[job_id][
                    "completed"
                ] += 1


                print(
                    f"Translated row "
                    f"{index + 1}/"
                    f"{total_rows}"
                )


            except Exception as error:

                print(
                    f"Translation error "
                    f"on row "
                    f"{index + 1}:",
                    error
                )


                # Keep empty if translation fails

                translations.append("")


                excel_jobs[job_id][
                    "failed"
                ] += 1


            # ----------------------------------
            # UPDATE PROGRESS
            # ----------------------------------

            excel_jobs[job_id][
                "current"
            ] = index + 1


        # --------------------------------------
        # ADD TRANSLATION COLUMN
        # --------------------------------------

        dataframe[
            target_column
        ] = translations


        # --------------------------------------
        # CREATE OUTPUT EXCEL
        # --------------------------------------

        output_file = BytesIO()


        with pd.ExcelWriter(
            output_file,
            engine="openpyxl"
        ) as writer:

            dataframe.to_excel(
                writer,

                index=False,

                sheet_name="Products"
            )


            worksheet = writer.book[
                "Products"
            ]


            format_worksheet(
                worksheet
            )


        output_file.seek(0)


        # --------------------------------------
        # STORE RESULT
        # --------------------------------------

        excel_jobs[job_id][
            "file"
        ] = output_file.getvalue()


        excel_jobs[job_id][
            "download_ready"
        ] = True


        excel_jobs[job_id][
            "status"
        ] = "completed"


        print(
            f"Excel job {job_id} completed."
        )


    except Exception as error:

        print(
            f"Excel job {job_id} failed:",
            error
        )


        excel_jobs[job_id][
            "status"
        ] = "failed"


        excel_jobs[job_id][
            "error"
        ] = str(error)


# ==========================================
# EXCEL TRANSLATION STATUS
# ==========================================

@app.route(
    "/translation-status/<job_id>",
    methods=["GET"]
)
def translation_status(job_id):

    if job_id not in excel_jobs:

        return jsonify({
            "error": "Job not found"
        }), 404


    job = excel_jobs[job_id]


    return jsonify({

        "status":
            job["status"],

        "current":
            job["current"],

        "total":
            job["total"],

        "completed":
            job["completed"],

        "failed":
            job["failed"],

        "download_ready":
            job["download_ready"],

        "error":
            job["error"]

    })


# ==========================================
# DOWNLOAD COMPLETED EXCEL
# ==========================================

@app.route(
    "/download-excel/<job_id>",
    methods=["GET"]
)
def download_excel(job_id):

    if job_id not in excel_jobs:

        return jsonify({
            "error":
                "Job not found"
        }), 404


    job = excel_jobs[job_id]


    if not job["download_ready"]:

        return jsonify({
            "error":
                "Excel file is not ready yet."
        }), 400


    output_file = BytesIO(
        job["file"]
    )


    output_file.seek(0)


    return send_file(
        output_file,

        as_attachment=True,

        download_name=(
            "translated_products.xlsx"
        ),

        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


# ==========================================
# START SERVER
# ==========================================

if __name__ == "__main__":

    app.run(
        debug=True,
        port=5000
    )